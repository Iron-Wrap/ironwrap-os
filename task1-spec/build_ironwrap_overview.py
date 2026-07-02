from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/ivan/Code/ironwrap")
OUT_DOC = ROOT / "output" / "documents"
ASSETS = ROOT / "task1-spec" / "assets"
DOCX = OUT_DOC / "IronWrap_OS_Overview.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(25, 25, 25)
MUTED = RGBColor(95, 95, 95)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "C8D0DA"


def font_path() -> str | None:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    return next((p for p in candidates if os.path.exists(p)), None)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = font_path()
    if path:
        return ImageFont.truetype(path, size=size, index=0)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        test = f"{line} {word}".strip()
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            line = test
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def draw_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str,
    fill: str,
    outline: str = "#355C7D",
) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=2)
    title_font = load_font(22, bold=True)
    body_font = load_font(18)
    draw.text((x1 + 18, y1 + 14), title, fill="#102A43", font=title_font)
    y = y1 + 50
    for line in wrap_text(draw, body, body_font, x2 - x1 - 36):
        draw.text((x1 + 18, y), line, fill="#243B53", font=body_font)
        y += 24


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line((start, end), fill="#54708F", width=4)
    ex, ey = end
    sx, sy = start
    if ex >= sx:
        points = [(ex, ey), (ex - 14, ey - 8), (ex - 14, ey + 8)]
    else:
        points = [(ex, ey), (ex + 14, ey - 8), (ex + 14, ey + 8)]
    draw.polygon(points, fill="#54708F")


def make_diagrams() -> tuple[Path, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    title_font = load_font(24, bold=True)
    small_font = load_font(17)

    first = Image.new("RGB", (1400, 720), "white")
    d = ImageDraw.Draw(first)
    d.text((36, 28), "First Version: Simple Operating Spine", fill="#102A43", font=title_font)
    boxes = {
        "lead": (40, 110, 315, 245),
        "ghl": (405, 110, 680, 245),
        "jn": (770, 110, 1045, 245),
        "work": (1135, 110, 1370, 245),
        "docs": (405, 345, 680, 480),
        "proposal": (770, 345, 1045, 480),
        "money": (1135, 345, 1370, 480),
        "ai": (40, 345, 315, 480),
    }
    draw_box(d, boxes["lead"], "Customer", "Website, phone, SMS, referral, ad", "#F7FAFC")
    draw_box(d, boxes["ghl"], "GHL", "Lead inbox, booking, SMS, reviews", "#EAF4FF")
    draw_box(d, boxes["jn"], "JobNimbus", "Job record, stages, tasks, owner", "#EAF7F0")
    draw_box(d, boxes["work"], "Work Flow", "Appointment, estimate, material request", "#FFF8E6")
    draw_box(d, boxes["ai"], "AI Pilot", "Missed calls and after-hours only", "#F4F0FF")
    draw_box(d, boxes["docs"], "Files", "One Drive/Google folder per job", "#F7FAFC")
    draw_box(d, boxes["proposal"], "Proposal", "Simple HTML/PDF package", "#F7FAFC")
    draw_box(d, boxes["money"], "QuickBooks", "Deposit, invoice, payment status", "#F7FAFC")
    for a, b in [
        ("lead", "ghl"),
        ("ghl", "jn"),
        ("jn", "work"),
        ("ai", "ghl"),
        ("jn", "docs"),
        ("docs", "proposal"),
        ("proposal", "money"),
    ]:
        ax = boxes[a][2]
        ay = (boxes[a][1] + boxes[a][3]) // 2
        bx = boxes[b][0]
        by = (boxes[b][1] + boxes[b][3]) // 2
        if a == "jn" and b == "docs":
            arrow(d, ((boxes[a][0] + boxes[a][2]) // 2, boxes[a][3]), ((boxes[b][0] + boxes[b][2]) // 2, boxes[b][1]))
        elif a == "ai":
            arrow(d, (boxes[a][2], ay), (boxes[b][0], by))
        else:
            arrow(d, (ax, ay), (bx, by))
    d.text((40, 620), "Rule: existing tools first; buy expensive add-ons only after volume proves the need.", fill="#52606D", font=small_font)
    first_path = ASSETS / "ironwrap_first_version_flow.png"
    first.save(first_path)

    end = Image.new("RGB", (1400, 760), "white")
    d = ImageDraw.Draw(end)
    d.text((36, 28), "End Version: Full IronWrap Operating System", fill="#102A43", font=title_font)
    boxes2 = {
        "entry": (40, 105, 315, 235),
        "front": (405, 105, 680, 235),
        "job": (770, 105, 1045, 235),
        "sales": (1135, 105, 1370, 235),
        "field": (40, 380, 315, 510),
        "shop": (405, 380, 680, 510),
        "finance": (770, 380, 1045, 510),
        "owner": (1135, 380, 1370, 510),
    }
    draw_box(d, boxes2["entry"], "Lead Sources", "Website, Renoworks, ads, phone, referrals", "#F7FAFC")
    draw_box(d, boxes2["front"], "Front Door", "GHL, AI receptionist, booking, nurture", "#EAF4FF")
    draw_box(d, boxes2["job"], "Job Core", "JobNimbus, folders, documents, tasks", "#EAF7F0")
    draw_box(d, boxes2["sales"], "Sales", "Visualizer, proposal, signature, deposit", "#FFF8E6")
    draw_box(d, boxes2["field"], "Field", "Mobile forms, photos, completion, issues", "#F4F0FF")
    draw_box(d, boxes2["shop"], "Materials", "Supplier/shop workflow, later Paradigm", "#F4F0FF")
    draw_box(d, boxes2["finance"], "Finance", "QuickBooks, payments, reminders", "#F7FAFC")
    draw_box(d, boxes2["owner"], "Owner View", "Daily exceptions, stuck jobs, revenue risk", "#EAF4FF")
    for a, b in [
        ("entry", "front"),
        ("front", "job"),
        ("job", "sales"),
        ("field", "shop"),
        ("shop", "finance"),
        ("finance", "owner"),
    ]:
        ax = boxes2[a][2]
        ay = (boxes2[a][1] + boxes2[a][3]) // 2
        bx = boxes2[b][0]
        by = (boxes2[b][1] + boxes2[b][3]) // 2
        arrow(d, (ax, ay), (bx, by))
    d.line(((905, 235), (905, 300), (175, 300), (175, 380)), fill="#54708F", width=4)
    d.polygon([(175, 380), (167, 366), (183, 366)], fill="#54708F")
    d.text((40, 300), "JobNimbus coordinates field work, materials, finance status, and owner exceptions.", fill="#52606D", font=small_font)
    d.text((40, 660), "Goal: every lead, job, document, payment, material request, and follow-up has one visible owner and next step.", fill="#52606D", font=small_font)
    end_path = ASSETS / "ironwrap_end_version_flow.png"
    end.save(end_path)
    return first_path, end_path


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Pt(widths[idx] / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, widths: list[int], header=True) -> None:
    set_table_width(table, widths)
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9.2)
            if header and r == 0:
                set_cell_shading(cell, LIGHT_GRAY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)


def set_doc_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for list_style in ["List Bullet", "List Number"]:
        st = styles[list_style]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.167


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = "IronWrap Operating System Overview"
    hp.style = doc.styles["Normal"]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in hp.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = MUTED
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("Confidential working overview")
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("IronWrap Operating System Overview")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("Simple non-technical build plan for the first version and the end-state system")
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED

    meta = doc.add_table(rows=4, cols=2)
    rows = [
        ("Audience", "IronWrap owner and operations team"),
        ("Purpose", "Show what should be built first, what comes later, and how to keep cost low"),
        ("Not a proposal", "This is a plain overview for deciding the operating system direction"),
        ("Date", "June 27, 2026"),
    ]
    for i, (k, v) in enumerate(rows):
        meta.rows[i].cells[0].text = k
        meta.rows[i].cells[1].text = v
    style_table(meta, [1900, 7460], header=False)
    for row in meta.rows:
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for p in row.cells[0].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = DARK_BLUE


def add_para(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        add_para(doc, item, "List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(f"{i}.  {item}")


def add_note(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF8E6")
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(122, 90, 0)
    p.add_run(body)
    style_table(table, [9360], header=False)


def add_decision_table(doc: Document) -> None:
    data = [
        ("Area", "Minimum / free-first", "Low-cost paid", "Buy later only if needed"),
        ("AI calls", "Human live-answer plus GHL missed-call workflow", "GHL AI for missed calls and after-hours", "Vapi, ElevenLabs, Retell, or custom voice only if GHL fails tests"),
        ("LLM helpers", "Manual templates and approved scripts", "OpenRouter free/cheap models for non-sensitive drafting and summaries", "Dedicated paid LLM only after clear usage exists"),
        ("OCR", "AppliCad XML, PDF text extraction, and manual review", "Tesseract for simple local OCR", "Mistral OCR 4 for hard PDFs, handwriting, or layout-heavy docs"),
        ("Storage", "One existing Google Drive or OneDrive structure", "Google Workspace/Microsoft 365 paid storage as needed", "Custom file storage only if folders become a bottleneck"),
        ("Visualizer", "Photos, product samples, and simple proposal links", "Renoworks later if visual selling becomes core", "3D/custom visualization only after sales proof"),
        ("Dashboard", "JobNimbus/GHL reports plus a simple exceptions sheet", "Looker Studio or small internal page", "Full dashboard only if owner still lacks daily visibility"),
        ("Shop/materials", "JobNimbus material request task and attachment", "Supplier forms/imports", "Paradigm only after manual material flow is stable"),
    ]
    table = doc.add_table(rows=len(data), cols=4)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = value
    style_table(table, [1450, 2450, 2450, 3010])


def add_cost_table(doc: Document) -> None:
    data = [
        ("Version", "Expected vendor cost", "What it means"),
        ("First version", "CAD $50-$400/mo new incremental if existing subscriptions are already paid; otherwise roughly CAD $1,100-$2,500/mo all-in", "Build the spine and avoid buying add-ons before volume proves need"),
        ("Expansion after proof", "CAD $4,200-$5,800/mo", "Add visualizer and predictable report volume only when sales process justifies it"),
        ("Full future system", "CAD $5,000-$6,800/mo plus known setup fees if Paradigm/Renoworks are purchased", "Only after materials, field, shop, and reporting volume justify a larger stack"),
    ]
    table = doc.add_table(rows=len(data), cols=3)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = value
    style_table(table, [1600, 3100, 4660])


def add_compliance_table(doc: Document) -> None:
    data = [
        ("Requirement", "What the OS must do"),
        ("Consent and notice", "Tell customers why their personal information is collected and use it only for that purpose."),
        ("Minimal data", "Collect only what IronWrap needs: contact, property, job, communications, appointment, proposal, invoice status, and documents."),
        ("Call recording and AI", "Disclose AI assistance and any recording/transcription at the start of the call, state the purpose, and offer a human path."),
        ("SMS/email follow-up", "For marketing, review requests, and nurture messages, store consent, identify IronWrap, and include unsubscribe/stop handling."),
        ("Outbound calls", "For telemarketing-style outbound calls, respect the National DNCL and IronWrap's internal do-not-call list."),
        ("Vendor processing", "Keep a vendor register, review terms, and tell customers if personal information may be handled by service providers outside Canada."),
        ("Access control", "Named users, MFA where available, least-privilege access, and no shared production passwords after testing."),
        ("Sensitive finance data", "Do not store credit card, banking, PAD, or financing documents in AI prompts, transcripts, broad CRM notes, or shared folders."),
        ("Retention and deletion", "Keep files/transcripts only as long as needed for the job, accounting, warranty, or legal purposes; delete or archive on schedule."),
        ("Breach response", "Keep breach records and notify the applicable commissioner/affected people when the legal risk threshold is met."),
    ]
    table = doc.add_table(rows=len(data), cols=2)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            table.rows[r].cells[c].text = value
    style_table(table, [2300, 7060])


def add_sources(doc: Document) -> None:
    add_para(doc, "Sources Checked", "Heading 1")
    add_para(
        doc,
        "These are included for traceability. The build should still receive owner approval and legal review before production launch.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(
        "Checked: OPC Canada PIPEDA and call-recording guidance; Alberta PIPA and OIPC breach guidance; "
        "CRTC CASL and National DNCL rules; OpenRouter pricing/free-model references; Google One storage; "
        "Tesseract OCR documentation; Mistral OCR 4 pricing and document-processing notes."
    )
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def build_doc() -> None:
    OUT_DOC.mkdir(parents=True, exist_ok=True)
    first_diagram, end_diagram = make_diagrams()

    doc = Document()
    doc.core_properties.title = "IronWrap Operating System Overview"
    doc.core_properties.subject = "Non-technical overview of the first and end-state IronWrap OS"
    doc.core_properties.author = "Mindburn Labs"
    doc.core_properties.comments = "Generated as a working overview for IronWrap."
    set_doc_styles(doc)
    add_header_footer(doc)
    add_title_block(doc)

    add_para(doc, "Recommendation In One Page", "Heading 1")
    add_note(
        doc,
        "Decision",
        "Build a simple operating spine first. Do not buy the full future stack until lead volume, report volume, and shop/material complexity prove it.",
    )
    add_para(
        doc,
        "IronWrap already has the core tools needed to start: GHL, JobNimbus, Make, folders, QuickBooks, and real measurement documents. The first version should connect and clean these tools, not replace them.",
    )
    add_bullets(
        doc,
        [
            "GHL should be the front door for leads, booking, SMS, review requests, and the first AI pilot.",
            "JobNimbus should be the main job record: customer, property, job number, stage, tasks, and files.",
            "Make should stay as automation glue, but every important automation needs clear logging and a manual fallback.",
            "QuickBooks should stay the accounting system. Do not let AI or CRM tools own final finance truth.",
            "AI should start with missed calls and after-hours calls, not replace human answering during business hours.",
        ],
    )

    add_para(doc, "First Version", "Heading 1")
    add_para(
        doc,
        "The first version should be small enough to ship and useful enough to change daily operations.",
    )
    add_numbered(
        doc,
        [
            "One lead inbox in GHL for website, phone, SMS, referrals, ads, and manual entry.",
            "Clean handoff from GHL to JobNimbus after a lead is qualified.",
            "JobNimbus owns the customer, property, job number, stage, task owner, and next action.",
            "One folder per real job, created only after the job number exists.",
            "GHL calendar for appointments and confirmations.",
            "AI pilot for missed calls and after-hours calls, with human handoff.",
            "Simple proposal package: HTML/customer view, PDF archive, product choices, payment/deposit instructions.",
            "Measurement/document intake: parse structured XML first, extract PDF text second, OCR only when needed.",
            "QuickBooks visibility for deposit, invoice, paid, overdue, and follow-up status.",
            "Review request and unhappy-customer branch after completion.",
            "Simple exceptions view: stuck leads, missed calls, overdue tasks, unpaid invoices, failed automations.",
        ],
    )
    doc.add_picture(str(first_diagram), width=Inches(6.3))

    add_para(doc, "What Not To Buy First", "Heading 1")
    add_bullets(
        doc,
        [
            "Do not buy Paradigm until manual material requests are stable.",
            "Do not buy NaturalForms unless JobNimbus mobile forms fail in real field work.",
            "Do not buy Renoworks unless visualizer-led selling is a deliberate sales strategy.",
            "Do not buy Vapi, ElevenLabs, Retell, or custom Twilio/OpenAI voice until GHL AI fails real call tests.",
            "Do not budget 60 RoofRuler reports per month until qualified job volume proves it.",
            "Do not build a custom dashboard until existing reports plus a simple exceptions sheet fail.",
        ],
    )

    add_para(doc, "End Version", "Heading 1")
    add_para(
        doc,
        "The end version is a full operating system for IronWrap, but it should be reached in steps. The goal is that every customer, job, document, material request, invoice, and follow-up has one visible owner and next step.",
    )
    add_bullets(
        doc,
        [
            "Website and Renoworks visualizer feed product interest directly into GHL and JobNimbus.",
            "AI receptionist handles after-hours, missed calls, basic qualification, booking, summaries, and safe human transfer.",
            "JobNimbus runs residential, commercial, and shop/supply workflows with the right fields and statuses.",
            "Documents, measurements, photos, and proposals are attached to the right job folder automatically.",
            "Field staff complete mobile checklists/forms and completion reports.",
            "Material requests become structured shop/supplier work, with Paradigm added only when justified.",
            "QuickBooks remains accounting truth, with payment reminders and status back to JobNimbus.",
            "Owner sees daily exceptions: slow leads, stuck jobs, missing deposits, material delays, failed automations, and review opportunities.",
        ],
    )
    doc.add_picture(str(end_diagram), width=Inches(6.3))

    add_para(doc, "Cost-Control Options", "Heading 1")
    add_para(
        doc,
        "The expensive parts should have gates. Each row starts with the cheapest path, then a low-cost path, then the full paid option only when the cheaper path fails.",
    )
    add_decision_table(doc)

    add_para(doc, "Budget View", "Heading 1")
    add_cost_table(doc)
    add_note(
        doc,
        "First budget to present",
        "If the current subscriptions are already paid, the first build should add very little new monthly spend. The main work is configuration, cleanup, automation discipline, and a controlled AI pilot.",
    )

    add_para(doc, "Canadian Compliance Guardrails", "Heading 1")
    add_para(
        doc,
        "The OS can be built in a compliant-by-design way for Canadian operations if these go-live gates are completed before production use. This is not a lawyer certification, but it is aligned to the official Canadian and Alberta privacy/communications requirements checked for this document.",
    )
    add_compliance_table(doc)
    add_note(
        doc,
        "Go-live rule",
        "No full production launch until consent language, call recording notice, unsubscribe handling, vendor register, access controls, retention rules, and breach response steps are approved by IronWrap.",
    )

    add_para(doc, "Practical First Build Sequence", "Heading 1")
    add_numbered(
        doc,
        [
            "Clean JobNimbus fields, stages, job numbering, and task ownership.",
            "Audit and simplify existing Make scenarios so one automation owns each action.",
            "Standardize GHL lead intake, dedupe, booking, and follow-up.",
            "Connect qualified leads to JobNimbus and folders without duplicates.",
            "Build the missed-call/after-hours AI pilot and test it with real call scripts.",
            "Create the simple proposal/PDF flow and attach it to the job folder.",
            "Add measurement/document intake with free extraction first and paid OCR only when needed.",
            "Turn on payment/review follow-ups and the simple exceptions view.",
            "Run test jobs end to end.",
            "Only then complete security/access hardening and production launch.",
        ],
    )

    add_sources(doc)
    doc.save(DOCX)


if __name__ == "__main__":
    build_doc()
    print(DOCX)
