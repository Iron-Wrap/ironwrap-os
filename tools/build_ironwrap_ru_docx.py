from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOC_SPECS = [
    (
        ROOT / "docs" / "archive" / "2026-06-27" / "IRONWRAP_TASK1_INTERNAL_OS_RU.md",
        ROOT / "output" / "documents" / "IronWrap_Task1_Internal_OS_RU.docx",
    ),
    (
        ROOT / "docs" / "IRONWRAP_TASK2_METAL_LAYOUT_PRODUCT_RU.md",
        ROOT / "output" / "documents" / "IronWrap_Task2_Metal_Layout_Product_RU.docx",
    ),
]

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
BLACK = RGBColor(0, 0, 0)
BULLET_NUM_ID = 101
DECIMAL_NUM_ID = 102


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Calibri", size=11, color=BLACK, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style._element.rPr.rFonts.set(qn("w:cs"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def set_paragraph_spacing(style, before=0, after=6, line=1.25):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_paragraph_border_bottom(paragraph, color="2E74B5", size="8", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_custom_numbering(doc):
    numbering = doc.part.numbering_part.element

    def add_level(abstract_id, num_id, num_fmt, lvl_text, left, hanging):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))

        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")

        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)

        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        lvl.append(fmt)

        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), lvl_text)
        lvl.append(text)

        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")
        lvl.append(suff)

        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        p_pr.append(tabs)

        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        p_pr.append(ind)
        lvl.append(p_pr)

        if num_fmt == "bullet":
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), "Symbol")
            r_fonts.set(qn("w:hAnsi"), "Symbol")
            r_pr.append(r_fonts)
            lvl.append(r_pr)

        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)

    add_level(101, BULLET_NUM_ID, "bullet", "•", 540, 270)
    add_level(102, DECIMAL_NUM_ID, "decimal", "%1.", 540, 360)


def set_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_inline_runs(paragraph, text, default_size=11):
    parts = re.split(r"(`[^`]+`|\\*\\*[^*]+\\*\\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=default_size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Courier New", size=default_size - 0.5, bold=False)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=default_size, bold=False)


def add_body_para(doc, text):
    p = doc.add_paragraph(style="Normal")
    add_inline_runs(p, text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    set_numbering(p, BULLET_NUM_ID)
    add_inline_runs(p, text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    set_numbering(p, DECIMAL_NUM_ID)
    add_inline_runs(p, text)
    return p


def configure_document(doc):
    add_custom_numbering(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    set_style_font(doc.styles["Normal"], size=11, color=BLACK)
    set_paragraph_spacing(doc.styles["Normal"], before=0, after=5, line=1.15)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[style_name]
        set_style_font(style, size=size, color=color, bold=True)
        set_paragraph_spacing(style, before=before, after=after, line=1.15)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        set_style_font(style, size=11, color=BLACK)
        set_paragraph_spacing(style, before=0, after=2, line=1.15)
        style.paragraph_format.left_indent = Inches(0.32)
        style.paragraph_format.first_line_indent = Inches(-0.16)

    section.header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""


def build_one(source, out):
    text = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    title = text[0].lstrip("# ").strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    set_run_font(run, size=24, color=BLACK, bold=True)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    set_paragraph_border_bottom(rule, color="2E74B5", size="8", space="6")

    pending_number = None
    in_quote = False
    for raw in text[1:]:
        line = raw.strip()
        if not line:
            in_quote = False
            continue

        if line.startswith("## "):
            heading_text = line[3:].strip()
            p = doc.add_paragraph(heading_text, style="Heading 1")
            keep_with_next(p)
            continue

        if line.startswith("**") and line.endswith("**  "):
            line = line[:-2].strip()

        if line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, line)
            keep_with_next(p)
            continue

        if line.startswith("- "):
            add_bullet(doc, line[2:].strip())
            continue

        numbered = re.match(r"^(\\d+)\\.\\s+(.*)$", line)
        if numbered:
            add_number(doc, numbered.group(2).strip())
            continue

        if line.startswith("> "):
            p = add_body_para(doc, line[2:].strip())
            set_shading(p, "F4F6F9")
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            for run in p.runs:
                run.font.italic = True
            in_quote = True
            continue

        if in_quote:
            p = add_body_para(doc, line)
            set_shading(p, "F4F6F9")
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.right_indent = Inches(0.15)
            for run in p.runs:
                run.font.italic = True
            continue

        p = add_body_para(doc, line)
        if line.startswith("Главный вывод:"):
            for run in p.runs:
                run.bold = True

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


def build():
    for source, out in DOC_SPECS:
        build_one(source, out)


if __name__ == "__main__":
    build()
